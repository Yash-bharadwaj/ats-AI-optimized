import PyPDF2
from docx import Document
import io
import logging
from typing import Optional

logger = logging.getLogger(__name__)

class FileProcessor:
    """Vercel-optimized file processor with minimal dependencies"""
    
    async def extract_text(self, content: bytes, file_extension: str) -> str:
        """Extract text from file content based on file type"""
        logger.info(f"Starting text extraction for file type: {file_extension}, size: {len(content)} bytes")
        
        try:
            if file_extension == 'pdf':
                logger.info("Processing PDF file...")
                return await self._extract_pdf_text(content)
            elif file_extension == 'docx':
                logger.info("Processing DOCX file...")
                return await self._extract_docx_text(content)
            elif file_extension == 'txt':
                logger.info("Processing text file...")
                return await self._extract_text_file(content)
            else:
                logger.error(f"Unsupported file type: {file_extension}")
                raise ValueError(f"Unsupported file type: {file_extension}. Only PDF, DOCX, and TXT files are supported.")
                
        except Exception as e:
            logger.error(f"Error extracting text from {file_extension} file: {str(e)}")
            raise Exception(f"Failed to extract text: {str(e)}")
    
    async def _extract_pdf_text(self, content: bytes) -> str:
        """Extract text from PDF using PyPDF2 (lightweight alternative)"""
        try:
            logger.info("Attempting PDF text extraction with PyPDF2...")
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
            text = ""
            
            for page_num in range(len(pdf_reader.pages)):
                page = pdf_reader.pages[page_num]
                page_text = page.extract_text()
                text += page_text + "\n"
                logger.debug(f"Extracted {len(page_text)} characters from page {page_num + 1}")
                
            if not text.strip():
                logger.warning("No text found in PDF file using PyPDF2")
                raise Exception("No text found in PDF file")
            
            logger.info(f"Successfully extracted {len(text)} characters from PDF")
            return text.strip()
            
        except Exception as e:
            logger.error(f"Error extracting PDF text with PyPDF2: {str(e)}")
            raise Exception(f"Failed to extract text from PDF: {str(e)}")
    
    async def _extract_docx_text(self, content: bytes) -> str:
        """Extract text from DOCX file"""
        try:
            logger.info("Attempting DOCX text extraction...")
            doc = Document(io.BytesIO(content))
            text = ""
            
            paragraph_count = 0
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
                    paragraph_count += 1
            
            # Also extract text from tables
            table_count = 0
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + " "
                    text += "\n"
                table_count += 1
            
            logger.info(f"Extracted text from {paragraph_count} paragraphs and {table_count} tables")
            
            if not text.strip():
                logger.warning("No text found in DOCX file")
                raise Exception("No text found in DOCX file")
            
            logger.info(f"Successfully extracted {len(text)} characters from DOCX")
            return text.strip()
            
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {str(e)}")
            raise Exception(f"Failed to extract text from DOCX: {str(e)}")
    
    async def _extract_text_file(self, content: bytes) -> str:
        """Extract text from plain text file"""
        try:
            logger.info("Processing text file...")
            # Try UTF-8 first
            text = content.decode('utf-8')
        except UnicodeDecodeError:
            try:
                # Try latin-1 as fallback
                text = content.decode('latin-1')
                logger.info("Used latin-1 encoding for text file")
            except UnicodeDecodeError:
                try:
                    # Try cp1252 as another fallback
                    text = content.decode('cp1252')
                    logger.info("Used cp1252 encoding for text file")
                except UnicodeDecodeError:
                    # Final fallback with error handling
                    text = content.decode('utf-8', errors='ignore')
                    logger.warning("Used UTF-8 with ignored errors for text file")
            
        logger.info(f"Extracted {len(text)} characters from text file")
        return text
    
    def validate_file_type(self, filename: str) -> bool:
        """Validate if file type is supported (Vercel-optimized)"""
        if not filename:
            return False
            
        file_extension = filename.split('.')[-1].lower()
        # Only support PDF, DOCX, and TXT for Vercel deployment
        supported_types = ['pdf', 'docx', 'txt']
        
        return file_extension in supported_types
    
    def validate_file_size(self, file_size: int, max_size: int = 4194304) -> bool:
        """Validate file size (default 4MB limit for Vercel)"""
        return file_size <= max_size 